"""H7.4/H7.5 Exporte und Berichte: Excel des Rasters, Portfolio-Bericht, Deckblatt-PDF.

Alles wird lokal erzeugt (openpyxl, PyMuPDF). Kennzahlen und Hinweise sind
deterministisch aus den Registerdaten berechnet, nicht von der KI.
"""

from __future__ import annotations

import html
import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database.models import Contract
from app.database.models.contract import ContractDocument
from app.database.models.deadlines import ContractDeadline
from app.database.models.document import DocumentFile, DocumentVersion
from app.database.models.register import ContractRemark
from app.finance.calc import month_add
from app.finance.service import finance_summary
from app.register.status import STATUS_LABELS, is_running
from app.reports.grid import COLUMN_BY_KEY, MONEY_KEYS, GridRow, GridSpec, MoneyTotals, groups, money_totals

_HEADER_FILL = PatternFill("solid", fgColor="DCE6F0")
_DATE_FORMAT = "DD.MM.YYYY"
_MONEY_FORMAT = '#,##0.00 "€"'


def _cell_value(column_key: str, value: Any) -> Any:
    column = COLUMN_BY_KEY[column_key]
    if value is None:
        return None
    if column.type == "status":
        return STATUS_LABELS.get(value, value)
    if column.type == "bool":
        return "Ja" if value else "Nein"
    if column.type == "money":
        return float(value)
    return value


def _write_sheet(sheet, headers: list[str], rows: list[list[Any]], formats: list[Optional[str]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for values in rows:
        sheet.append([None] * len(values))
        row_index = sheet.max_row
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=index)
            cell.value = value
            if isinstance(value, str):
                # Keine Formeln aus Vertragsdaten (Formel-Injektion).
                cell.data_type = "s"
            if formats[index - 1] and value is not None:
                cell.number_format = formats[index - 1]
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(row[index - 1])) for row in rows[:500] if row[index - 1] is not None])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _float_or_none(value: Optional[str]) -> Optional[float]:
    return float(Decimal(value)) if value is not None else None


def _coverage_text(summary: dict, key: str) -> str:
    coverage = summary["coverage"][key]
    total = coverage["known"] + coverage["missing"]
    if not total:
        return ""
    if coverage["missing"] == 0:
        return f"vollständig ({total} Verträge)"
    return f"bekannt für {coverage['known']} von {total} Verträgen"


def _format_for(column_key: str) -> Optional[str]:
    kind = COLUMN_BY_KEY[column_key].type
    return _DATE_FORMAT if kind == "date" else _MONEY_FORMAT if kind == "money" else None


def grid_xlsx(rows: list[GridRow], spec: GridSpec, *, title: str, as_of: date) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Verträge"
    columns = list(dict.fromkeys([*spec.group_by, *spec.columns]))
    _write_sheet(
        sheet,
        [COLUMN_BY_KEY[key].label for key in columns],
        [[_cell_value(key, row.values[key]) for key in columns] for row in rows],
        [_format_for(key) for key in columns],
    )
    if spec.group_by:
        group_sheet = workbook.create_sheet("Gruppen")
        group_rows = [
            [" / ".join(item["labels"]), item["count"], *[_float_or_none(item["sums"][key]) for key in MONEY_KEYS]]
            for item in groups(rows, spec)
        ]
        _write_sheet(
            group_sheet,
            ["Gruppe", "Anzahl", *[COLUMN_BY_KEY[key].label for key in MONEY_KEYS]],
            group_rows,
            [None, None, *[_MONEY_FORMAT for _ in MONEY_KEYS]],
        )
    info = workbook.create_sheet("Info")
    summary = money_totals(rows).payload()
    info.append(["Bericht", title])
    info.append(["Stand", as_of.strftime("%d.%m.%Y")])
    info.append(["Anzahl Verträge", summary["count"]])
    for key in MONEY_KEYS:
        info.append([f"Summe {COLUMN_BY_KEY[key].label}", _float_or_none(summary["sums"][key]), _coverage_text(summary, key)])
        info.cell(row=info.max_row, column=2).number_format = _MONEY_FORMAT
        for currency, amount in summary["other_currencies"].get(key, {}).items():
            info.append([f"Summe {COLUMN_BY_KEY[key].label} ({currency})", float(Decimal(amount))])
    info.append(["Hinweis", "Werte deterministisch aus den Stammdaten und Finanzdaten berechnet (netto, EUR). "
                 "Leere Summen bedeuten: keine gepflegten Werte, nicht 0."])
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 60
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- portfolio report

_DIMENSIONS = (
    ("hospital", "Krankenhaus"),
    ("department", "Abteilung"),
    ("cost_center", "Kostenstelle"),
    ("contract_type", "Vertragstyp"),
    ("partner", "Vertragspartner"),
)
TOP_N = 10


def _money(value: Optional[Decimal]) -> Optional[str]:
    return str(value.quantize(Decimal("0.01"))) if value is not None else None


def _amount(totals_: MoneyTotals, key: str) -> dict:
    """Betrag mit Abdeckung fuer Kennzahlen (None = nichts bekannt, nicht 0)."""
    payload = totals_.payload()
    return {
        "value": payload["sums"][key],
        "currency": payload["currency"],
        "known": payload["coverage"][key]["known"],
        "missing": payload["coverage"][key]["missing"],
        "other_currencies": payload["other_currencies"].get(key, {}),
    }


def running_rows(rows: list[GridRow], as_of: date) -> list[GridRow]:
    return [
        row for row in rows
        if is_running(row.values["business_status"], row.values["term_start_date"], row.values["possible_end"], as_of)
    ]


def _merge(parts: list[MoneyTotals]) -> MoneyTotals:
    merged = MoneyTotals()
    for part in parts:
        merged.count += part.count
        for key in MONEY_KEYS:
            merged.sums[key] += part.sums[key]
            merged.main_known[key] += part.main_known[key]
            merged.known[key] += part.known[key]
            merged.missing[key] += part.missing[key]
    return merged


def _bucket_item(label: str, value: Optional[str], data: MoneyTotals) -> dict:
    return {
        "label": label, "value": value, "count": data.count,
        "annual_cost": _money(data.sum("annual_cost")), "cost_current_year": _money(data.sum("cost_current_year")),
        "known": data.known["cost_current_year"], "missing": data.missing["cost_current_year"],
    }


NONE_VALUE = "__none__"
_DRILL_LEVELS = (("hospital", "Krankenhaus"), ("department", "Abteilung"))
DRILL_CONTRACT_LIMIT = 200


def _drill_match(value: Optional[str], wanted: str) -> bool:
    return (value or None) == (None if wanted == NONE_VALUE else wanted)


def drilldown(rows: list[GridRow], as_of: date, hospital: Optional[str] = None, department: Optional[str] = None) -> dict:
    """Analysen: Organisation -> Krankenhaus -> Abteilung -> Verträge (laufende Verträge, Summen mit Abdeckung).

    `hospital`/`department` sind die angezeigten Namen; NONE_VALUE steht für „ohne Zuordnung“.
    """
    if department is not None and hospital is None:
        raise ValueError("Abteilung nur innerhalb eines Krankenhauses.")
    running = running_rows([row for row in rows if row.values["business_status"] != "geloescht"], as_of)
    path: list[dict] = []
    register_filter: list[dict] = []
    for (key, label), wanted in zip(_DRILL_LEVELS, (hospital, department)):
        if wanted is None:
            break
        running = [row for row in running if _drill_match(row.values[key], wanted)]
        path.append({"level": key, "label": label, "value": wanted, "name": "Ohne Zuordnung" if wanted == NONE_VALUE else wanted})
        register_filter.append({"column": key, "op": "in", "value": ["" if wanted == NONE_VALUE else wanted]})

    level = _DRILL_LEVELS[len(path)][0] if len(path) < len(_DRILL_LEVELS) else "contract"
    if level == "contract":
        ordered = sorted(
            running,
            key=lambda row: (row.values["cost_current_year"] is None, -(row.values["cost_current_year"] or 0), row.values["register_number"] or ""),
        )
        items = [{
            "contract_id": str(row.contract_id),
            "register_number": row.values["register_number"],
            "short_description": row.values["short_description"],
            "partner": row.values["partner"],
            "currency": row.values.get("_currency") or "EUR",
            "annual_cost": _money(row.values["annual_cost"]),
            "cost_current_year": _money(row.values["cost_current_year"]),
            "term_status": row.values["term_status"],
        } for row in ordered[:DRILL_CONTRACT_LIMIT]]
    else:
        buckets: dict[Optional[str], MoneyTotals] = defaultdict(MoneyTotals)
        for row in running:
            buckets[row.values[level] or None].add(row)
        ordered_buckets = sorted(
            buckets.items(),
            key=lambda item: (-(item[1].sum("cost_current_year") or Decimal("-1")), item[0] is None, item[0] or ""),
        )
        items = [_bucket_item(name or "Ohne Zuordnung", name or NONE_VALUE, data) for name, data in ordered_buckets]

    return {
        "as_of": as_of.isoformat(),
        "level": level,
        "path": path,
        "items": items,
        "count": len(running),
        "truncated": level == "contract" and len(running) > DRILL_CONTRACT_LIMIT,
        "totals": _amount(money_totals(running), "cost_current_year"),
        "register_filter": register_filter,
    }


def portfolio(rows: list[GridRow], as_of: date, years: range) -> dict:
    active = [row for row in rows if row.values["business_status"] != "geloescht"]
    running = running_rows(active, as_of)
    running_ids = {row.contract_id for row in running}
    in_progress = [row for row in active if row.values["business_status"] == "in_arbeit"]

    by_status: dict[str, int] = defaultdict(int)
    for row in active:
        by_status[row.values["business_status"]] += 1

    cost_by: dict[str, list[dict]] = {}
    for key, _label in _DIMENSIONS:
        buckets: dict[Optional[str], MoneyTotals] = defaultdict(MoneyTotals)
        for row in running:
            # Kostenstellen: Betraege anteilig nach der gepflegten Aufteilung.
            shares = (
                [(name, share / Decimal("100")) for name, share in row.allocations]
                if key == "cost_center" and row.allocations else [(row.values[key], Decimal("1"))]
            )
            for name, share in shares:
                buckets[name].add(row, share)
        ordered = sorted(
            buckets.items(), key=lambda item: item[1].sum("cost_current_year") or Decimal("-1"), reverse=True
        )
        head, tail = ordered[:TOP_N], ordered[TOP_N:]
        items = [_bucket_item(name or "(ohne)", name or "", data) for name, data in head]
        if tail:
            items.append(_bucket_item(f"Übrige ({len(tail)})", None, _merge([data for _, data in tail])))
        cost_by[key] = items

    yearly_rows = running + [
        row for row in active
        if row.values["business_status"] in ("beendet", "gekuendigt") and row.contract_id not in running_ids
    ]
    yearly = {year: {"expense": Decimal("0"), "income": Decimal("0")} for year in years}
    with_finance = 0
    for row in yearly_rows:
        if row.yearly:
            with_finance += 1
        for year, values in row.yearly.items():
            if year in yearly:
                yearly[year]["expense"] += values["expense"]
                yearly[year]["income"] += values["income"]

    horizon = month_add(date(as_of.year, as_of.month, 1), 6)
    months = [month_add(date(as_of.year, as_of.month, 1), offset) for offset in range(6)]
    upcoming = []
    monthly: dict[str, dict[str, int]] = {month.strftime("%Y-%m"): {"notice": 0, "other": 0} for month in months}
    for row in running + in_progress:
        for due, title in row.upcoming:
            if as_of <= due < horizon:
                kind = "notice" if title == "Kündigungsfrist" or "Kündigung" in title else "other"
                monthly[due.strftime("%Y-%m")][kind] += 1
                upcoming.append({
                    "contract_id": str(row.contract_id), "register_number": row.values["register_number"],
                    "short_description": row.values["short_description"], "due_date": due.isoformat(), "title": title,
                    "kind": kind,
                })
    upcoming.sort(key=lambda item: item["due_date"])

    running_totals = money_totals(running)
    partner_top = cost_by["partner"][0] if cost_by["partner"] else None
    current_total = running_totals.sum("cost_current_year")
    concentration = (
        (Decimal(partner_top["cost_current_year"]) / current_total * 100).quantize(Decimal("0.1"))
        if partner_top and partner_top["value"] and partner_top["cost_current_year"] and current_total and current_total > 0
        else None
    )
    soon = [row for row in running if row.values["cancel_by"] and as_of <= row.values["cancel_by"] <= as_of.fromordinal(as_of.toordinal() + 90)]
    considered = running + in_progress
    insights = [
        {"key": "no_end", "label": "Laufzeitangaben unvollständig – Ende und Fristen nicht berechenbar",
         "count": sum(1 for row in considered if row.values["term_status"] == "Unvollständig"),
         "filter": {"column": "term_status", "op": "in", "value": ["Unvollständig"]}, "severity": "critical"},
        {"key": "no_responsible", "label": "Laufende Verträge ohne verantwortliche Person",
         "count": sum(1 for row in running if not row.values["responsible"]),
         "filter": {"column": "responsible", "op": "empty"}, "severity": "warning"},
        {"key": "notice_no_reminder", "label": "Kündigungsfrist in 90 Tagen ohne Erinnerungsempfänger",
         "count": sum(1 for row in soon if not row.values["reminder_recipient"]),
         "filter": {"column": "cancel_by", "op": "within_days", "value": 90},
         # Genau die betroffenen Verträge (Frist bald UND kein Empfänger); `filter` bleibt für ältere Aufrufer.
         "filters": [{"column": "cancel_by", "op": "within_days", "value": 90}, {"column": "reminder_recipient", "op": "empty"}],
         "severity": "warning"},
        # Auch "In Arbeit": sonst zeigt die Kachel 0, solange noch kein Vertrag laeuft.
        {"key": "no_finance", "label": "Verträge ohne gepflegte Kosten (laufend oder in Arbeit)",
         "count": sum(1 for row in considered if row.values["annual_cost"] is None and row.values["cost_current_year"] is None),
         "filter": {"column": "cost_current_year", "op": "empty"}, "severity": "warning"},
        {"key": "deletion_due", "label": "Aufbewahrungsfrist abgelaufen (Löschprüfung)",
         "count": sum(1 for row in active if row.values["deletion_from"] and row.values["deletion_from"] <= as_of),
         "filter": {"column": "deletion_from", "op": "overdue"}, "severity": "warning"},
    ]

    return {
        "as_of": as_of.isoformat(),
        "definitions": {
            "running": "Gültig oder gekündigt und am Stichtag innerhalb der Laufzeit (Beginn erreicht, Ende nicht überschritten). "
                       "Verträge „In Arbeit“ zählen gesondert.",
            "annual_cost": "Jährliche Kosten (aktueller Satz, netto) aus den gepflegten Finanzposten laufender Verträge. "
                           "KI-Vorschläge sind nicht enthalten.",
        },
        "totals": {
            "contracts": len(active),
            "running": len(running),
            "in_progress": len(in_progress),
            "perpetual": sum(1 for row in considered if row.values["term_status"] == "Unbefristet"),
            "annual_cost": _money(running_totals.sum("annual_cost")),
            "cost_current_year": _money(current_total),
            "income_current_year": _money(running_totals.sum("income_current_year")),
            "amounts": {
                key: _amount(running_totals, key) for key in ("annual_cost", "cost_current_year", "income_current_year")
            },
        },
        "by_status": [
            {"status": status, "label": STATUS_LABELS.get(status, status), "count": by_status.get(status, 0)}
            for status in ("in_arbeit", "gueltig", "gekuendigt", "beendet")
        ],
        "cost_by": cost_by,
        "dimensions": [{"key": key, "label": label} for key, label in _DIMENSIONS],
        "cost_by_year": [
            {"year": year, "expense": _money(values["expense"]) if with_finance else None,
             "income": _money(values["income"]) if with_finance else None, "is_projection": year > as_of.year}
            for year, values in yearly.items()
        ],
        "cost_by_year_coverage": {"known": with_finance, "total": len(yearly_rows)},
        "deadlines_by_month": [{"month": key, **values} for key, values in monthly.items()],
        "upcoming": upcoming[:100],
        "supplier_concentration": (
            {"partner": partner_top["label"], "share_percent": str(concentration)} if concentration is not None else None
        ),
        "insights": insights,
    }


def _amount_coverage(amount: dict) -> str:
    total = amount["known"] + amount["missing"]
    if not total:
        return "keine laufenden Verträge"
    if amount["missing"] == 0:
        return f"vollständig ({total} Verträge)"
    return f"bekannt für {amount['known']} von {total} Verträgen"


def portfolio_xlsx(report: dict) -> bytes:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Übersicht"
    totals_ = report["totals"]
    _write_sheet(
        overview, ["Kennzahl", "Wert"],
        [
            ["Stand", date.fromisoformat(report["as_of"])],
            ["Verträge (ohne gelöschte)", totals_["contracts"]],
            ["Davon laufend", totals_["running"]],
            ["Jährliche Kosten (netto)", _float_or_none(totals_["annual_cost"])],
            ["Kosten lfd. Jahr (netto)", _float_or_none(totals_["cost_current_year"])],
            ["Einnahmen lfd. Jahr (netto)", _float_or_none(totals_["income_current_year"])],
            ["In Arbeit (nicht laufend)", totals_.get("in_progress", 0)],
            ["Abdeckung jährliche Kosten", _amount_coverage(totals_["amounts"]["annual_cost"])],
            *[[item["label"], item["count"]] for item in report["by_status"]],
            *[[item["label"], item["count"]] for item in report["insights"]],
        ],
        [None, None],
    )
    for row_index in (2,):
        overview.cell(row=row_index, column=2).number_format = _DATE_FORMAT
    for row_index in (5, 6, 7):
        overview.cell(row=row_index, column=2).number_format = _MONEY_FORMAT
    for dimension in report["dimensions"]:
        sheet = workbook.create_sheet(f"Kosten je {dimension['label']}"[:31])
        _write_sheet(
            sheet, [dimension["label"], "Anzahl", "Jährliche Kosten (netto)", "Kosten lfd. Jahr (netto)"],
            [[item["label"], item["count"], _float_or_none(item["annual_cost"]), _float_or_none(item["cost_current_year"])]
             for item in report["cost_by"][dimension["key"]]],
            [None, None, _MONEY_FORMAT, _MONEY_FORMAT],
        )
    years = workbook.create_sheet("Kosten je Jahr")
    _write_sheet(
        years, ["Jahr", "Ausgaben (netto)", "Einnahmen (netto)", "Hochrechnung"],
        [[item["year"], _float_or_none(item["expense"]), _float_or_none(item["income"]), "Ja" if item["is_projection"] else "Nein"]
         for item in report["cost_by_year"]],
        [None, _MONEY_FORMAT, _MONEY_FORMAT, None],
    )
    upcoming = workbook.create_sheet("Fristen 6 Monate")
    _write_sheet(
        upcoming, ["Fällig", "Nr.", "Vertrag", "Termin"],
        [[date.fromisoformat(item["due_date"]), item["register_number"], item["short_description"], item["title"]]
         for item in report["upcoming"]],
        [_DATE_FORMAT, None, None, None],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- cover sheet PDF

_CSS = """
body { font-family: sans-serif; font-size: 9.5pt; color: #1d2733; }
h1 { font-size: 16pt; margin: 0 0 2pt 0; }
h2 { font-size: 11pt; margin: 12pt 0 4pt 0; border-bottom: 1px solid #9aa7b5; }
p.meta { color: #55606d; margin: 0 0 6pt 0; }
table { width: 100%; border-collapse: collapse; }
td, th { padding: 2pt 4pt; vertical-align: top; text-align: left; }
th { width: 34%; color: #55606d; font-weight: normal; }
tr.head th, tr.head td { font-weight: bold; color: #1d2733; border-bottom: 1px solid #c8d0da; }
td.num { text-align: right; }
p.note { color: #55606d; font-size: 8pt; }
"""


def _fmt(value: Any, kind: str = "text") -> str:
    if value is None or value == "":
        return "–"
    if kind == "date" and isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if kind == "money":
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        text = f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"{text} €"
    if kind == "bool":
        return "Ja" if value else "Nein"
    if kind == "status":
        return STATUS_LABELS.get(value, str(value))
    return str(value)


def _rows(pairs: list[tuple[str, str]]) -> str:
    return "".join(f"<tr><th>{html.escape(label)}</th><td>{html.escape(value)}</td></tr>" for label, value in pairs)


def cover_html(session: Session, contract: Contract, row: GridRow, as_of: date, organization_name: str) -> str:
    values = row.values

    def field(key: str) -> tuple[str, str]:
        column = COLUMN_BY_KEY[key]
        return column.label, _fmt(values[key], column.type)

    sections = [
        ("Vertrag & Partner", ["short_description", "business_status", "partner", "contract_type", "contract_category",
                               "internal_number", "is_framework", "data_protection_relevant", "is_confidential"]),
        ("Laufzeit & Fristen", ["term_model", "term_start_date", "earliest_end", "cancel_by", "possible_end",
                                "terminated_on", "terminated_effective", "next_deadline"]),
        ("Verantwortung", ["responsible", "reminder_recipient", "reminder_date", "escalation_recipient"]),
        ("Organisation & Ablage", ["hospital", "department", "cost_center", "storage_location", "deletion_from"]),
    ]
    parts = [
        f"<h1>Vertrag {html.escape(values['register_number'] or 'ohne Nummer')}</h1>",
        f"<p class='meta'>{html.escape(organization_name)} · Deckblatt · Stand {as_of.strftime('%d.%m.%Y')}</p>",
    ]
    for title, keys in sections:
        parts.append(f"<h2>{html.escape(title)}</h2><table>{_rows([field(key) for key in keys])}</table>")
    if contract.subject_text:
        parts.append(f"<h2>Vertragsgegenstand</h2><p>{html.escape(contract.subject_text)}</p>")

    summaries = finance_summary(session, contract, as_of)
    parts.append("<h2>Finanzen (netto)</h2>")
    if not summaries:
        parts.append("<p>Keine Finanzdaten gepflegt.</p>")
    for summary in summaries:
        currency = summary["currency"]
        lines = [("", "Ausgaben", "Einnahmen")]
        for key, label in (("annual_run_rate", "Jährlich (aktueller Satz)"), ("current_year", f"Laufendes Jahr {as_of.year}"),
                           ("to_date", "Bis heute"), ("remaining_to_end", "Bis Vertragsende"), ("total_term", "Gesamtlaufzeit")):
            expense = (summary["expense"].get(key) or {}).get("net") if summary["expense"].get(key) else None
            income = (summary["income"].get(key) or {}).get("net") if summary["income"].get(key) else None
            lines.append((label, _fmt(expense, "money") if expense is not None else "–", _fmt(income, "money") if income is not None else "–"))
        table = "".join(
            f"<tr class='{'head' if index == 0 else ''}'><th>{html.escape(a)}</th><td class='num'>{html.escape(b)}</td><td class='num'>{html.escape(c)}</td></tr>"
            for index, (a, b, c) in enumerate(lines)
        )
        parts.append(f"<p class='meta'>Währung {html.escape(currency)}</p><table>{table}</table>")
        for note in summary.get("notes") or []:
            parts.append(f"<p class='note'>{html.escape(note)}</p>")

    deadlines = session.scalars(
        select(ContractDeadline).where(ContractDeadline.contract_id == contract.id, ContractDeadline.status == "open")
        .order_by(ContractDeadline.due_date).limit(15)
    ).all()
    parts.append("<h2>Offene Termine</h2>")
    parts.append(
        "<table>" + _rows([(_fmt(item.due_date, "date"), item.title) for item in deadlines]) + "</table>"
        if deadlines else "<p>Keine offenen Termine.</p>"
    )

    documents = session.execute(
        select(ContractDocument, DocumentFile.original_filename)
        .join(DocumentVersion, DocumentVersion.id == ContractDocument.document_version_id)
        .join(DocumentFile, DocumentFile.id == DocumentVersion.document_file_id)
        .where(ContractDocument.contract_id == contract.id)
        .order_by(ContractDocument.is_main.desc(), ContractDocument.created_at)
    ).all()
    parts.append("<h2>Dokumente</h2>")
    parts.append(
        "<table>" + _rows([
            ("Hauptdokument" if link.is_main else (link.attachment_type or "Anhang"),
             f"{link.title or filename}{'' if link.is_valid else ' (ungültig)'}")
            for link, filename in documents
        ]) + "</table>"
        if documents else "<p>Keine Dokumente.</p>"
    )

    remarks = session.scalars(
        select(ContractRemark).where(ContractRemark.contract_id == contract.id).order_by(ContractRemark.created_at.desc()).limit(3)
    ).all()
    if remarks:
        parts.append("<h2>Letzte Bemerkungen</h2><table>" + _rows([
            (f"{_fmt(item.created_at.date(), 'date')} · {item.author_name or ''}", item.body[:400]) for item in remarks
        ]) + "</table>")
    parts.append("<p class='note'>Fristen und Beträge wurden deterministisch aus den Stammdaten berechnet. Erstellt lokal mit MedFlux.</p>")
    return "<html><body>" + "".join(parts) + "</body></html>"


def html_to_pdf(document_html: str) -> bytes:
    import fitz

    story = fitz.Story(html=document_html, user_css=_CSS)
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (48, 48, -48, -48)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()
