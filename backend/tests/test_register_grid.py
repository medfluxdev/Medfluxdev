"""H7 Vertragsregister: Raster, Filter, Gruppen, Ansichten, Excel, Portfolio, Deckblatt."""

from __future__ import annotations

import io
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from types import SimpleNamespace
from uuid import UUID, uuid4

import fitz
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routes import contract_finance, contract_register, register_grid
from app.core.scope_context import ScopeContext
from app.database import scope_queries
from app.database.models import Account, AccountMembership, Contract
from app.database.models.organization import Organization
from app.database.models.organization_structure import Hospital
from app.database.session import create_session
from app.reports.grid import GridError, GridSpec


def test_spec_parsing_is_strict():
    assert GridSpec.parse(None).columns[:2] == ["register_number", "short_description"]
    for bad in (
        {"filters": [{"column": "nope", "op": "in"}]},
        {"filters": [{"column": "partner", "op": "drop"}]},
        {"group_by": ["short_description"]},
        {"sort": {"column": "nope"}},
    ):
        try:
            GridSpec.parse(bad)
        except GridError:
            continue
        raise AssertionError(bad)


def _seed():
    ids = {key: uuid4() for key in ("org", "hospital_a", "hospital_b", "admin", "manager", "viewer")}
    slug = f"grid-{uuid4().hex[:12]}"
    session = create_session()
    try:
        session.add(Organization(id=ids["org"], name="Raster Klinik", slug=slug))
        session.add_all([
            Hospital(id=ids["hospital_a"], organization_id=ids["org"], name="Klinikum A", code="A"),
            Hospital(id=ids["hospital_b"], organization_id=ids["org"], name="Klinikum B", code="B"),
        ])
        for key in ("admin", "manager", "viewer"):
            session.add(Account(id=ids[key], login_name=f"{key}-{slug}@x.invalid", display_name=key.title(), credential_digest="t"))
        session.flush()
        for key, role in (("admin", "admin"), ("manager", "contract_manager"), ("viewer", "viewer")):
            session.add(AccountMembership(organization_id=ids["org"], account_id=ids[key], role=role))
        session.commit()
    finally:
        session.close()
    return slug, ids


def _as(monkeypatch, ids, key, role):
    actor = SimpleNamespace(organization_id=ids["org"], account_id=ids[key], role=role, is_developer=False, display_name=key.title())
    for module in (contract_register, register_grid):
        monkeypatch.setattr(module, "current_actor", lambda request, organization_slug: actor)
        monkeypatch.setattr(module, "current_scope_context", lambda: ScopeContext())
    monkeypatch.setattr(scope_queries, "current_identity", lambda: actor)
    app = FastAPI()
    for module in (contract_register, contract_finance, register_grid):
        app.include_router(module.router)
    return TestClient(app)


def _populate(client, slug, ids):
    query = {"organization_slug": slug}
    a = client.post("/api/contracts", params=query, json={"hospital_id": str(ids["hospital_a"]), "values": {
        "short_description": "Wartung MRT", "term_model": "fixed_term", "term_start_date": "2026-01-01",
        "term_fixed_end_date": "2027-12-31",
    }}).json()["contract_id"]
    b = client.post("/api/contracts", params=query, json={"hospital_id": str(ids["hospital_b"]), "values": {
        "short_description": "=SUM(1)", "term_model": "terminable", "term_start_date": "2020-01-01",
        "notice_period_value": 1, "notice_period_unit": "months", "notice_to": "month_end",
    }}).json()["contract_id"]
    c = client.post("/api/contracts", params=query, json={"values": {"short_description": "Altvertrag"}}).json()["contract_id"]
    for cid, amount, interval in ((a, "1000", "monthly"), (b, "6000", "annual")):
        response = client.post(f"/api/contracts/{cid}/finance/entries", params=query, json={"values": {
            "label": "Pauschale", "amount": amount, "vat_rate": "19", "interval": interval, "valid_from": "2026-01-01",
        }})
        assert response.status_code == 200, response.text
    session = create_session()
    try:
        session.get(Contract, UUID(a)).responsible_account_id = ids["manager"]
        session.get(Contract, UUID(b)).responsible_account_id = None
        session.get(Contract, UUID(c)).business_status = "geloescht"
        session.commit()
    finally:
        session.close()
    return a, b, c


def _grid(client, slug, spec=None, **params):
    response = client.get("/api/register/grid", params={"organization_slug": slug, "spec": json.dumps(spec or {}), **params})
    assert response.status_code == 200, response.text
    return response.json()


def test_grid_filters_groups_sort_and_tree(monkeypatch):
    slug, ids = _seed()
    admin = _as(monkeypatch, ids, "admin", "admin")
    a, b, c = _populate(admin, slug, ids)

    columns = admin.get("/api/register/columns", params={"organization_slug": slug}).json()["columns"]
    assert {"cancel_by", "annual_cost", "deletion_from", "hospital"} <= {item["key"] for item in columns}

    grid = _grid(admin, slug)
    assert grid["total"] == 2
    assert grid["totals"]["sums"]["annual_cost"] == "18000.00"
    assert grid["items"][0]["values"]["business_status"] == "In Arbeit"

    assert _grid(admin, slug, {"filters": [{"column": "hospital", "op": "in", "value": ["Klinikum A"]}]})["total"] == 1
    assert _grid(admin, slug, {"filters": [{"column": "cancel_by", "op": "within_days", "value": 90}]})["items"][0]["contract_id"] == b
    assert _grid(admin, slug, {"filters": [{"column": "annual_cost", "op": "range", "value": {"from": "10000"}}]})["total"] == 1
    assert _grid(admin, slug, {"filters": [{"column": "business_status", "op": "in", "value": ["geloescht"]}]})["items"][0]["contract_id"] == c
    assert _grid(admin, slug, {"include_deleted": True})["total"] == 3

    ordered = _grid(admin, slug, {"sort": {"column": "annual_cost", "dir": "desc"}})
    assert [item["contract_id"] for item in ordered["items"]] == [a, b]

    grouped = _grid(admin, slug, {"group_by": ["hospital"]})
    assert [(item["labels"], item["count"], item["sums"]["annual_cost"]) for item in grouped["groups"]] == [
        (["Klinikum A"], 1, "12000.00"), (["Klinikum B"], 1, "6000.00"),
    ]
    assert admin.get("/api/register/grid", params={"organization_slug": slug, "spec": "{bad"}).status_code == 422

    values = admin.get("/api/register/grid/values", params={"organization_slug": slug, "column": "hospital",
                                                            "spec": json.dumps({"filters": [{"column": "hospital", "op": "in", "value": ["Klinikum A"]}]})}).json()
    assert [item["label"] for item in values["values"]] == ["Klinikum A", "Klinikum B"]

    tree = admin.get("/api/register/grid/tree", params={"organization_slug": slug}).json()
    assert [(item["label"], item["count"]) for item in tree["hospitals"]] == [("Klinikum A", 1), ("Klinikum B", 1)]

    manager = _as(monkeypatch, ids, "manager", "contract_manager")
    mine = _grid(manager, slug, {"filters": [{"column": "responsible", "op": "me"}]})
    assert [item["contract_id"] for item in mine["items"]] == [a]


def test_archive_view_shows_only_archived_contracts(monkeypatch):
    slug, ids = _seed()
    admin = _as(monkeypatch, ids, "admin", "admin")
    a, b, _ = _populate(admin, slug, ids)
    session = create_session()
    try:
        row = session.get(Contract, UUID(a))
        row.is_archived = True
        row.archived_at = datetime(2026, 9, 30, 8, 0, tzinfo=timezone.utc)
        session.commit()
    finally:
        session.close()

    active = _grid(admin, slug, {"include_deleted": True})
    assert a not in [item["contract_id"] for item in active["items"]]
    assert b in [item["contract_id"] for item in active["items"]]

    archived = _grid(admin, slug, {"archived": True, "columns": ["register_number", "archived_at"]})
    assert [item["contract_id"] for item in archived["items"]] == [a]
    assert archived["spec"]["archived"] is True
    assert archived["items"][0]["values"]["archived_at"] == "2026-09-30"


def test_saved_views_defaults_permissions_and_ownership(monkeypatch):
    slug, ids = _seed()
    query = {"organization_slug": slug}
    viewer = _as(monkeypatch, ids, "viewer", "viewer")
    listed = viewer.get("/api/register/views", params=query).json()
    names = [item["name"] for item in listed["items"]]
    assert names == ["Meine Verträge", "Meine Wiedervorlage", "Kündigungsfrist < 90 Tage", "Zu prüfen", "Regelverstöße", "Archiv", "In Arbeit", "Zum Löschen vorgesehen", "Gelöschte Verträge"]
    assert {item["system_key"] for item in listed["items"]} >= {"review", "archive"}
    assert listed["can_manage_global"] is False
    system_id = listed["items"][0]["id"]

    assert viewer.post("/api/register/views", params=query, json={"name": "Global", "is_global": True}).status_code == 403
    own = viewer.post("/api/register/views", params=query, json={"name": "Meine CT", "config": {"q": "CT"}})
    assert own.status_code == 200 and own.json()["is_own"] is True
    assert viewer.post("/api/register/views", params=query, json={"name": "x", "config": {"filters": [{"column": "x", "op": "in"}]}}).status_code == 422
    assert viewer.put(f"/api/register/views/{system_id}", params=query, json={"name": "Umbenannt"}).status_code == 403

    manager = _as(monkeypatch, ids, "manager", "contract_manager")
    assert "Meine CT" not in [item["name"] for item in manager.get("/api/register/views", params=query).json()["items"]]
    assert manager.delete(f"/api/register/views/{own.json()['id']}", params=query).status_code == 404

    admin = _as(monkeypatch, ids, "admin", "admin")
    assert admin.delete(f"/api/register/views/{system_id}", params=query).status_code == 403
    shared = admin.post("/api/register/views", params=query, json={"name": "Für alle", "is_global": True, "config": {"group_by": ["hospital"]}})
    assert shared.json()["is_global"] is True
    assert admin.delete(f"/api/register/views/{shared.json()['id']}", params=query).json() == {"deleted": True}

    viewer = _as(monkeypatch, ids, "viewer", "viewer")
    assert viewer.delete(f"/api/register/views/{own.json()['id']}", params=query).status_code == 200


def test_excel_export_portfolio_and_cover(monkeypatch):
    slug, ids = _seed()
    query = {"organization_slug": slug}
    admin = _as(monkeypatch, ids, "admin", "admin")
    a, b, _ = _populate(admin, slug, ids)

    exported = admin.get("/api/register/grid/export.xlsx", params={**query, "spec": json.dumps({"group_by": ["hospital"], "columns": ["register_number", "short_description", "annual_cost", "cancel_by"]})})
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content))
    sheet = workbook["Verträge"]
    assert [cell.value for cell in sheet[1]] == ["Gesellschaft", "Nr.", "Kurzbeschreibung", "Jährliche Kosten (netto)", "Kündbar bis"]
    injected = next(row for row in sheet.iter_rows(min_row=2) if row[2].value == "=SUM(1)")
    assert injected[2].data_type == "s"
    assert injected[4].number_format == "DD.MM.YYYY"
    assert workbook["Gruppen"]["A2"].value == "Klinikum A"

    report = admin.get("/api/register/reports/portfolio", params=query).json()
    # Neu angelegte Verträge sind "In Arbeit" und zählen nicht als laufend.
    assert report["totals"]["running"] == 0 and report["totals"]["in_progress"] == 2
    assert report["totals"]["annual_cost"] is None
    assert report["totals"]["amounts"]["annual_cost"] == {
        "value": None, "currency": "EUR", "known": 0, "missing": 0, "other_currencies": {},
    }
    session = create_session()
    try:
        for cid in (a, b):
            session.get(Contract, UUID(cid)).business_status = "gueltig"
        session.commit()
    finally:
        session.close()
    report = admin.get("/api/register/reports/portfolio", params=query).json()
    assert report["totals"]["running"] == 2 and report["totals"]["annual_cost"] == "18000.00"
    assert report["totals"]["amounts"]["annual_cost"]["known"] == 2
    assert report["totals"]["amounts"]["annual_cost"]["missing"] == 0
    assert report["cost_by"]["hospital"][0]["label"] == "Klinikum A"
    year = next(item for item in report["cost_by_year"] if item["year"] == date.today().year)
    assert Decimal(year["expense"]) > 0
    assert any(item["key"] == "no_responsible" and item["count"] == 1 for item in report["insights"])
    assert admin.get("/api/register/reports/portfolio.xlsx", params=query).status_code == 200

    pdf = admin.get(f"/api/contracts/{a}/cover.pdf", params=query)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    text = "".join(page.get_text() for page in fitz.open(stream=pdf.content, filetype="pdf"))
    assert "Deckblatt" in text and "Wartung MRT" in text and "Kündbar bis" in text and "12.000,00 €" in text

    viewer = _as(monkeypatch, ids, "viewer", "viewer")
    assert viewer.get("/api/register/grid/export.xlsx", params=query).status_code == 403
    assert viewer.get("/api/register/reports/portfolio", params=query).status_code == 200


def test_viewers_may_manage_only_their_own_views_through_middleware():
    from app.auth.middleware import _self_service_mutation

    assert _self_service_mutation("/api/register/views")
    assert _self_service_mutation("/api/register/views/0193a6c2-0000-7000-8000-000000000001")
    assert not _self_service_mutation("/api/register/views/../grid")
    assert not _self_service_mutation("/api/register/grid/export.xlsx")
